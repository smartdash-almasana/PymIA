from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from pymia.smartpyme.service_1_xlsx_to_normalized_table_v1 import (
    read_xlsx_to_normalized_tables_v1,
)


def _write(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet(title=name)
        first = False
        sheet.title = name
        for row in rows:
            sheet.append(list(row))
    workbook.save(path)
    return path


def _sucursales_rows() -> list[list[object]]:
    return [
        ["SucursalID", "Sucursal", "Ciudad"],
        ["S001", "Centro", "Queretaro"],
        ["S002", "Juriquilla", "Queretaro"],
        ["S003", "Antea", "Queretaro"],
        ["S004", "Roma Norte", "CDMX"],
        ["S005", "Polanco", "CDMX"],
    ]


def _by_sheet(tables: list[dict]) -> dict[str, dict]:
    return {str(table.get("sheet_name")): table for table in tables if table.get("status") == "OK"}


def test_all_text_header_with_digit_id_rows_is_accepted(tmp_path: Path) -> None:
    path = _write(tmp_path / "tie.xlsx", {"Sucursales": _sucursales_rows()})
    tables = read_xlsx_to_normalized_tables_v1(path)
    by_sheet = _by_sheet(tables)
    assert "Sucursales" in by_sheet
    table = by_sheet["Sucursales"]
    assert table["headers"] == ["SucursalID", "Sucursal", "Ciudad"]
    assert list(table["rows"][0].keys()) == list(table["normalized_headers"])
    assert table["rows"][0][table["normalized_headers"][0]] == "S001"
    assert table["rows"][0][table["normalized_headers"][1]] == "Centro"
    assert len(table["rows"]) == 5


def test_same_structure_other_sheet_name_resolves_generically(tmp_path: Path) -> None:
    path = _write(tmp_path / "tie2.xlsx", {"Almacenes": _sucursales_rows()})
    by_sheet = _by_sheet(read_xlsx_to_normalized_tables_v1(path))
    assert "Almacenes" in by_sheet
    assert by_sheet["Almacenes"]["headers"] == ["SucursalID", "Sucursal", "Ciudad"]


def test_digit_symmetric_tie_keeps_veto(tmp_path: Path) -> None:
    path = _write(
        tmp_path / "ambiguous.xlsx",
        {
            "Ventas": [["Fecha", "Total"], ["2026-01-01", 100]],
            "Notas": [
                ["Informe", "General", "X"],
                ["uno", "dos", "tres"],
                ["cuatro", "cinco", "seis"],
            ],
        },
    )
    by_sheet = _by_sheet(read_xlsx_to_normalized_tables_v1(path))
    assert "Ventas" in by_sheet
    assert "Notas" not in by_sheet


def test_valid_numeric_table_still_passes(tmp_path: Path) -> None:
    path = _write(
        tmp_path / "valid.xlsx",
        {"Ventas": [["Fecha", "Total"], ["2026-01-01", 100], ["2026-01-02", 200]]},
    )
    by_sheet = _by_sheet(read_xlsx_to_normalized_tables_v1(path))
    assert by_sheet["Ventas"]["headers"] == ["Fecha", "Total"]
    assert len(by_sheet["Ventas"]["rows"]) == 2


def test_truly_empty_sheet_still_skipped(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vacia"
    path = tmp_path / "empty.xlsx"
    workbook.save(path)
    tables = read_xlsx_to_normalized_tables_v1(path)
    assert all(table.get("status") != "OK" for table in tables)


def test_header_detection_is_deterministic(tmp_path: Path) -> None:
    path = _write(tmp_path / "det.xlsx", {"Sucursales": _sucursales_rows()})
    first = read_xlsx_to_normalized_tables_v1(path)[0]
    second = read_xlsx_to_normalized_tables_v1(path)[0]
    assert first["headers"] == second["headers"]
    assert first["normalized_headers"] == second["normalized_headers"]
    assert first["rows"] == second["rows"]


def test_no_semantic_inference_in_tables(tmp_path: Path) -> None:
    path = _write(tmp_path / "nosem.xlsx", {"Sucursales": _sucursales_rows()})
    table = read_xlsx_to_normalized_tables_v1(path)[0]
    blob = str(table)
    for token in ("semantic_role", "semantic_meaning", "inferred_meaning", "Sucursales"):
        if token == "Sucursales":
            continue
        assert token not in blob
    assert table["sheet_name"] == "Sucursales"
