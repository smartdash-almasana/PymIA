from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel-prueba"
OUT.mkdir(parents=True, exist_ok=True)


def save(name: str, title: str, headers: list[str], rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(OUT / name)


# 001 — mixed currency: same semantic amount column carries ARS and USD without FX evidence.
save(
    "S1_A4_ADV_001_mixed_currency.xlsx",
    "Ventas",
    ["fecha", "comprobante", "moneda", "venta_total"],
    [
        ["2026-07-01", "F001", "ARS", 120000],
        ["2026-07-02", "F002", "USD", 150],
        ["2026-07-03", "F003", "ARS", 90000],
    ],
)

# 002 — subtotal row mixed with detail rows.
save(
    "S1_A4_ADV_002_subtotal_as_operation.xlsx",
    "Ventas",
    ["fecha", "comprobante", "producto", "cantidad", "venta_total"],
    [
        ["2026-07-01", "F001", "A", 1, 10000],
        ["2026-07-01", "F001", "B", 2, 20000],
        [None, "SUBTOTAL", None, None, 30000],
        ["2026-07-02", "F002", "C", 1, 15000],
        [None, "TOTAL", None, None, 45000],
    ],
)

# 003 — zero is valid evidence; blank means absent evidence.
save(
    "S1_A4_ADV_003_zero_vs_blank.xlsx",
    "Resumen",
    ["ventas_periodo", "cobros_periodo", "impuestos_periodo"],
    [
        [100000, 0, None],
    ],
)

# 004 — sign convention not declared.
save(
    "S1_A4_ADV_004_inverted_signs.xlsx",
    "Caja",
    ["fecha", "concepto", "importe"],
    [
        ["2026-07-01", "venta contado", -50000],
        ["2026-07-02", "pago proveedor", 30000],
        ["2026-07-03", "cobro cliente", -20000],
    ],
)

# 005 — rows outside the apparent analysis period.
save(
    "S1_A4_ADV_005_out_of_period_dates.xlsx",
    "Ventas",
    ["fecha", "venta_total"],
    [
        ["2025-12-31", 50000],
        ["2026-07-01", 10000],
        ["2026-07-15", 12000],
        ["2026-08-01", 25000],
    ],
)

# 006 — exact duplicate business rows, potentially legitimate or accidental.
save(
    "S1_A4_ADV_006_duplicate_rows.xlsx",
    "Ventas",
    ["fecha", "comprobante", "producto", "cantidad", "venta_total"],
    [
        ["2026-07-01", "F001", "A", 1, 10000],
        ["2026-07-01", "F001", "A", 1, 10000],
        ["2026-07-02", "F002", "B", 1, 15000],
    ],
)

# 007 — mixed granularity: invoice header totals plus line-level detail in one table.
save(
    "S1_A4_ADV_007_mixed_granularity.xlsx",
    "Ventas",
    ["fecha", "comprobante", "producto", "cantidad", "precio_unitario", "importe_factura"],
    [
        ["2026-07-01", "F001", "A", 1, 10000, 30000],
        ["2026-07-01", "F001", "B", 2, 10000, 30000],
        ["2026-07-02", "F002", "C", 1, 15000, 15000],
    ],
)

# 008 — materially missing input for net margin: taxes absent.
save(
    "S1_A4_ADV_008_missing_material_input.xlsx",
    "Resumen",
    ["ventas_periodo", "cmv_total"],
    [
        [100000, 60000],
    ],
)

# 009 — extreme but finite values; should not be auto-corrected merely for looking odd.
save(
    "S1_A4_ADV_009_extreme_values.xlsx",
    "Resumen",
    ["ventas_periodo", "cmv_total", "impuestos_periodo"],
    [
        [999999999999.0, 1.0, 0.0],
    ],
)

# 010 — semantic decoys: similarly named amount fields with distinct meanings.
save(
    "S1_A4_ADV_010_semantic_decoy_columns.xlsx",
    "Ventas",
    ["fecha", "importe", "importe_bruto", "importe_neto", "importe_cobrado", "total"],
    [
        ["2026-07-01", 10000, 12000, 9000, 8500, 10000],
        ["2026-07-02", 20000, 24000, 18000, 17000, 20000],
    ],
)

# 011 — incomplete relationship: sales and collections cannot be linked reliably.
wb = Workbook()
ventas = wb.active
ventas.title = "Ventas"
ventas.append(["venta_id", "fecha", "cliente", "venta_total"])
ventas.append(["V001", "2026-07-01", "Cliente A", 10000])
ventas.append(["V002", "2026-07-02", "Cliente B", 20000])
cobros = wb.create_sheet("Cobros")
cobros.append(["cobro_id", "fecha", "importe_cobrado", "referencia"])
cobros.append(["C001", "2026-07-03", 9000, None])
cobros.append(["C002", "2026-07-04", 15000, "REF-X"])
wb.save(OUT / "S1_A4_ADV_011_incomplete_relationships.xlsx")

print("\n".join(str(path.name) for path in sorted(OUT.glob("S1_A4_ADV_*.xlsx"))))
